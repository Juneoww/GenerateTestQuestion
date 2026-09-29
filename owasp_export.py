"""功能：
  将 OWASP 红队提示词的最终六字段记录导出为一张供测试人员直接使用的 Excel 工作表。
实现：
  导出前复用生成器的严格契约校验；只创建“测试提示词”工作表，固定六列、冻结首行并为
  长提示词开启自动换行；模型文本强制保存为字面字符串，避免 Excel 公式解释。
  不写入富用例、步骤、执行记录或隐藏元数据。
输入：
  已生成的六字段记录列表与目标 .xlsx 文件路径。
输出：
  写好的 .xlsx 路径；无效记录以 ValueError 拒绝。
依赖：
  openpyxl 和 owasp_prompt_generator。
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from owasp_prompt_generator import validate_prompt_records


EXCEL_SHEET_NAME = "测试提示词"
EXCEL_COLUMNS = ("风险编号", "风险名称", "userA", "userB", "prompt", "断言")
_FIELD_BY_COLUMN = {
    "风险编号": "riskId",
    "风险名称": "riskName",
    "userA": "userA",
    "userB": "userB",
    "prompt": "prompt",
    "断言": "assertion",
}
_COLUMN_WIDTHS = (14, 22, 42, 42, 48, 32)


def export_prompt_xlsx(records: list[dict[str, str]], path: Path) -> Path:
    """将严格校验后的记录写成唯一的六列工作表。"""
    errors = validate_prompt_records(records)
    if errors:
        raise ValueError(f"OWASP 提示词不符合六字段契约：{'；'.join(errors)}")

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = EXCEL_SHEET_NAME
    header_font = Font(bold=True)
    body_alignment = Alignment(wrap_text=True, vertical="top")

    for index, column in enumerate(EXCEL_COLUMNS, 1):
        cell = sheet.cell(row=1, column=index, value=column)
        cell.font = header_font
        sheet.column_dimensions[get_column_letter(index)].width = _COLUMN_WIDTHS[index - 1]

    for record in records:
        sheet.append([record[_FIELD_BY_COLUMN[column]] for column in EXCEL_COLUMNS])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            # openpyxl 默认把“=”开头的模型文本识别为公式；这里保留原文并强制写成字符串。
            cell.data_type = "s"
            cell.alignment = body_alignment
    sheet.freeze_panes = "A2"
    workbook.save(destination)
    return destination
