"""集成测试：导出的题库工作簿列顺序、题型标签和文本换行。"""
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import excel_export


class ExcelExportTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory(prefix="gtq_export_")
        self.addCleanup(self.temp_dir.cleanup)
        self.path = Path(self.temp_dir.name) / "questions.xlsx"

    def _load_export(self, questions: list[dict]):
        exported = excel_export.export_xlsx(questions, {"batchId": "batch-1"}, self.path)
        self.assertEqual(exported, self.path)
        workbook = load_workbook(self.path)
        self.addCleanup(workbook.close)
        return workbook["测试题库"]

    def test_export_keeps_workbook_basics_and_question_type_header(self):
        sheet = self._load_export([{"question": "题干"}])

        self.assertEqual(sheet.title, "测试题库")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.cell(row=1, column=3).value, "题型")

    def test_export_labels_image_text_and_legacy_questions(self):
        sheet = self._load_export([
            {"question": "图片题", "questionType": "image"},
            {"question": "文本题", "questionType": "text"},
            {"question": "旧版题"},
        ])

        self.assertEqual(sheet.cell(row=2, column=3).value, "图片")
        self.assertEqual(sheet.cell(row=3, column=3).value, "文本")
        self.assertEqual(sheet.cell(row=4, column=3).value, "文本")

    def test_export_wraps_evidence_excerpt_in_column_ten_for_data_rows(self):
        evidence_texts = ["第一段原文摘录", "第二段原文摘录"]
        sheet = self._load_export([
            {"question": "题一", "evidenceText": evidence_texts[0]},
            {"question": "题二", "evidenceText": evidence_texts[1]},
        ])

        for row, evidence_text in enumerate(evidence_texts, start=2):
            self.assertEqual(sheet.cell(row=row, column=10).value, evidence_text)
            self.assertTrue(sheet.cell(row=row, column=10).alignment.wrap_text)


if __name__ == "__main__":
    unittest.main()
